"""Turn what the user gives us (SMILES, a name, files, a CSV/Excel list) into ``MoleculeInput`` records."""

from __future__ import annotations

import csv
import json
import logging
import re
import urllib.parse
import urllib.request
from collections import Counter
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass, field, replace
from pathlib import Path

from rdkit import Chem, rdBase
from rdkit.Chem import rdDetermineBonds

_log = logging.getLogger(__name__)

# Column names are matched case-insensitively. They include the EPA CompTox Chemicals Dashboard export
# headers (DTXSID, PREFERRED_NAME, CASRN, SMILES, QSAR_READY_SMILES, MS_READY_SMILES).
SMILES_COLUMNS = ("smiles", "qsar_ready_smiles", "ms_ready_smiles", "canonical_smiles", "isomeric_smiles")
ID_COLUMNS = ("dtxsid", "id", "name", "compound", "preferred_name", "cas", "casrn")
NAME_COLUMNS = ("preferred_name", "name", "chemical_name", "compound")
TABLE_TYPES = {".csv", ".xlsx", ".xlsm", ".xls"}
FILE_TYPES = {".sdf", ".mol", ".xyz", ".pdb", ".smi", ".smiles", ".txt"} | TABLE_TYPES

# Which SMILES column to use when a table has several, in order of preference. A row whose preferred
# cell is blank (or does not parse in RDKit) falls back to the next column that has a structure.
#   "smiles": the registered structure as drawn, with salts, stereo and charges intact. The pipeline
#             does its own desalting, neutralization and pH protonation, so this keeps the most
#             information; stereocenters matter for 3D (e.g. the C2 stereocenter of HFPO-DA).
#   "qsar":   QSAR-ready SMILES: desalted, neutralized, stereo and isotopes removed, tautomers
#             normalized (Mansouri et al., J. Cheminform. 2024, 16, 19). Use it to match descriptor
#             tables built on QSAR-ready forms. Without stereo, RDKit embeds an arbitrary stereoisomer.
#   "ms":     MS-ready SMILES: QSAR-ready-like forms with mixtures split into components, made for
#             mass-spectrometry screening (McEachran et al., J. Cheminform. 2018, 10, 45).
SMILES_PREFERENCES: dict[str, tuple[str, ...]] = {
    "smiles": SMILES_COLUMNS,
    "qsar": ("qsar_ready_smiles", "smiles", "ms_ready_smiles", "canonical_smiles", "isomeric_smiles"),
    "ms": ("ms_ready_smiles", "qsar_ready_smiles", "smiles", "canonical_smiles", "isomeric_smiles"),
}
# Placeholders that CompTox exports and spreadsheets use for "no value".
MISSING_VALUES = {"", "n/a", "na", "nan", "none", "null", "-", "#n/a"}


@dataclass
class MoleculeInput:
    id: str
    smiles: str = ""
    name: str = ""
    source: str = ""
    mol: Chem.Mol | None = None      # set when read from a structure file
    charge: int | None = None        # None = take from the structure after protonation
    multiplicity: int = 1
    meta: dict = field(default_factory=dict)

    def to_mol(self) -> Chem.Mol:
        if self.mol is not None:
            return Chem.Mol(self.mol)
        mol = Chem.MolFromSmiles(self.smiles)
        if mol is None:
            raise ValueError(f"Could not parse SMILES for {self.id}: {self.smiles!r}")
        return mol


def safe_id(text: str, fallback: str = "mol") -> str:
    """Folder- and file-safe identifier."""
    s = re.sub(r"[^A-Za-z0-9._+-]+", "_", text.strip()).strip("._")
    return (s or fallback)[:80]


def looks_like_smiles(text: str) -> bool:
    text = text.strip()
    if not text or " " in text:
        return False
    with rdBase.BlockLogs():         # names are tried as SMILES first; keep RDKit's parse errors quiet
        if any(ch in text for ch in "=#()[]@+\\/") or text.isupper() and len(text) <= 3:
            return Chem.MolFromSmiles(text, sanitize=False) is not None
        # All-letter SMILES (CCCO, OCCO, NCCN, ClCCl, CCCCCCCCCCCCO) parse and sanitize; plain words such as
        # "benzene" or "ethanol" do not (lowercase letters outside b/c/n/o/p/s, or aromatic atoms
        # without a ring), so a sanitizing parse tells them apart.
        return Chem.MolFromSmiles(text) is not None


def from_smiles(smiles: str, id: str | None = None, name: str = "") -> MoleculeInput:
    smiles = smiles.strip()
    if Chem.MolFromSmiles(smiles) is None:
        raise ValueError(f"Not a valid SMILES string: {smiles!r}")
    return MoleculeInput(id=safe_id(id or name or smiles), smiles=smiles, name=name or id or "", source="smiles")


def lookup_name(name: str, timeout: float = 15.0) -> MoleculeInput:
    """Resolve a chemical name, CAS number or DTXSID-free identifier through PubChem."""
    quoted = urllib.parse.quote(name.strip())
    url = f"https://pubchem.ncbi.nlm.nih.gov/rest/pug/compound/name/{quoted}/property/SMILES,IsomericSMILES,Title/JSON"
    req = urllib.request.Request(url, headers={"User-Agent": "QuantaBench"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            data = json.load(resp)
    except Exception as exc:
        raise ValueError(f"PubChem lookup failed for {name!r}: {exc}") from exc
    props = (data.get("PropertyTable", {}).get("Properties") or [{}])[0]
    smiles = props.get("SMILES") or props.get("IsomericSMILES") or ""
    if not smiles:
        raise ValueError(f"PubChem returned no structure for {name!r}")
    title = props.get("Title") or name
    return MoleculeInput(id=safe_id(name), smiles=smiles, name=title, source="pubchem",
                         meta={"pubchem_cid": props.get("CID")})


def parse_text(text: str, skip_errors: bool = False, log: Callable[[str], None] | None = None,
               source: str = "") -> list[MoleculeInput]:
    """Each non-empty line is a SMILES (optionally followed by an id) or a name to look up.

    By default the first line that cannot be resolved raises. With ``skip_errors`` (list files), such
    lines are reported through ``log`` and skipped, as the CSV/Excel readers skip bad rows; it still
    raises when no line could be read at all."""
    out = []
    errors = []
    for lineno, line in enumerate(text.splitlines(), 1):
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        first, _, rest = line.partition(" ")
        try:
            if looks_like_smiles(first):
                out.append(from_smiles(first, id=rest.strip() or None, name=rest.strip()))
            elif any(ch in first for ch in "=#()[]@\\/") and not rest:
                raise ValueError(f"Not a valid SMILES string: {first!r}")
            else:
                out.append(lookup_name(line))
        except ValueError as exc:
            if not skip_errors:
                raise
            errors.append(f"{source + ' ' if source else ''}line {lineno}: {exc}; skipped")
    for e in errors:
        (log or _log.warning)(e)
    if errors and not out:
        raise ValueError(f"No molecule could be read{' from ' + source if source else ''}: " + "; ".join(errors[:5])
                         + (f" (and {len(errors) - 5} more)" if len(errors) > 5 else ""))
    return out


def _pick(columns: list[str], wanted: tuple[str, ...], explicit: str | None) -> str | None:
    if explicit:
        if explicit in columns:
            return explicit
        lower = {c.lower().strip(): c for c in columns}
        if explicit.lower().strip() in lower:  # forgiving about case: headers get retyped in Excel
            return lower[explicit.lower().strip()]
        raise ValueError(f"Column {explicit!r} not in file (columns: {', '.join(columns)})")
    lower = {c.lower().strip(): c for c in columns}
    for w in wanted:
        if w in lower:
            return lower[w]
    return None


def _pick_all(columns: list[str], wanted: Sequence[str]) -> list[str]:
    lower = {c.lower().strip(): c for c in columns}
    return [lower[w] for w in wanted if w in lower]


def _cell(value) -> str:
    """Spreadsheet or CSV cell -> clean string ('' for blanks and N/A placeholders)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel stores every number as a float
    text = str(value).strip()
    return "" if text.lower() in MISSING_VALUES else text


def _parses(smiles: str) -> bool:
    with rdBase.BlockLogs():  # thousands of rows: keep RDKit's parse errors out of the console
        return Chem.MolFromSmiles(smiles) is not None


def smiles_preference_columns(preference: str | Sequence[str] | None) -> tuple[str, ...]:
    """Resolve a preset ("smiles", "qsar", "ms"), a single column name (tried first, then the
    defaults) or an explicit ordered list of column names. Matching is case-insensitive."""
    if preference is None:
        return SMILES_PREFERENCES["smiles"]
    if isinstance(preference, str):
        key = preference.lower().strip()
        if key in SMILES_PREFERENCES:
            return SMILES_PREFERENCES[key]
        return (key, *(c for c in SMILES_COLUMNS if c != key))
    return tuple(c.lower().strip() for c in preference)


@dataclass
class ReadReport:
    """What a table reader found, so callers can tell the user what was skipped and why."""

    inputs: list[MoleculeInput] = field(default_factory=list)
    n_rows: int = 0                                                  # non-empty data rows
    skipped: list[tuple[str, str]] = field(default_factory=list)    # (row label, reason)
    smiles_columns: Counter = field(default_factory=Counter)        # column -> rows that used it
    warnings: list[str] = field(default_factory=list)

    def emit(self, log: Callable[[str], None] | None = None) -> None:
        for w in self.warnings:
            (log or _log.warning)(w)


def dedupe_ids(items: Iterable[MoleculeInput], warnings: list[str] | None = None) -> list[MoleculeInput]:
    """Make ids unique, because the id names the output folder and the job record.

    An exact repeat (same id and structure) is dropped; the same id with a different structure gets a
    ``_2``, ``_3``... suffix. The result depends only on input order, so re-reading the same list gives
    the same ids, which resumable runs rely on.
    """
    out: list[MoleculeInput] = []
    seen: dict[str, str] = {}
    dropped = 0
    for it in items:
        key = it.smiles or (Chem.MolToSmiles(it.mol) if it.mol is not None else "")
        if it.id in seen:
            if seen[it.id] == key:
                dropped += 1
                continue
            n = 2
            while f"{it.id}_{n}" in seen:
                n += 1
            if warnings is not None:
                warnings.append(f"Duplicate id {it.id!r} with a different structure renamed to {it.id}_{n}")
            it = replace(it, id=f"{it.id}_{n}")
        seen[it.id] = key
        out.append(it)
    if dropped and warnings is not None:
        warnings.append(f"Dropped {dropped} exact duplicate row(s) (same id and structure)")
    return out


def _rows_to_inputs(columns: list[str], rows: Iterable[tuple[int, int, dict]], file_name: str, stem: str,
                    smiles_col: str | None, id_col: str | None,
                    smiles_preference: str | Sequence[str] | None) -> ReadReport:
    """Shared CSV/Excel logic. ``rows`` yields (spreadsheet row number, data row index, {column: value})."""
    if smiles_col:
        scols = [_pick(columns, (), smiles_col)]
    else:
        scols = _pick_all(columns, smiles_preference_columns(smiles_preference))
    if not scols:
        raise ValueError(f"No SMILES column found in {file_name} (columns: {', '.join(columns)})")
    icol = _pick(columns, ID_COLUMNS, id_col)
    ncol = _pick(columns, NAME_COLUMNS, None)
    rep = ReadReport()
    fallbacks = 0
    for rownum, index, row in rows:
        cells = {k: _cell(v) for k, v in row.items() if k}
        if not any(cells.values()):
            continue  # fully empty line (Excel sheets often end with some)
        rep.n_rows += 1
        ident = cells.get(icol, "") if icol else ""
        candidates = [(c, cells.get(c, "")) for c in scols if cells.get(c, "")]
        if not candidates:
            rep.skipped.append((ident or f"row {rownum}", "no structure"))
            continue
        # the first column whose SMILES parses; if none parses keep the first, so the run reports it
        used, smi = next(((c, v) for c, v in candidates if _parses(v)), candidates[0])
        if used != scols[0]:
            fallbacks += 1
        rep.smiles_columns[used] += 1
        name = (cells.get(ncol, "") if ncol else "") or ident
        rep.inputs.append(MoleculeInput(id=safe_id(ident or f"{stem}_{index}"), smiles=smi, name=name,
                                        source=f"{file_name}:{rownum}",
                                        meta={**cells, "smiles_column": used, "row": rownum}))
    if rep.skipped:
        shown = ", ".join(label for label, _ in rep.skipped[:5]) + (", ..." if len(rep.skipped) > 5 else "")
        rep.warnings.append(f"{file_name}: skipped {len(rep.skipped)} row(s) without a structure ({shown})")
    if fallbacks:
        used = ", ".join(f"{c} for {n}" for c, n in rep.smiles_columns.items() if c != scols[0])
        rep.warnings.append(f"{file_name}: {fallbacks} row(s) had no usable {scols[0]}; used {used}")
    rep.inputs = dedupe_ids(rep.inputs, rep.warnings)
    return rep


def read_csv_report(path: Path, smiles_col: str | None = None, id_col: str | None = None,
                    smiles_preference: str | Sequence[str] | None = None) -> ReadReport:
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        columns = list(reader.fieldnames or [])
        rows = ((i + 1, i, row) for i, row in enumerate(reader, 1))  # header is line 1
        return _rows_to_inputs(columns, rows, path.name, path.stem, smiles_col, id_col, smiles_preference)


def read_csv(path: Path, smiles_col: str | None = None, id_col: str | None = None,
             smiles_preference: str | Sequence[str] | None = None,
             log: Callable[[str], None] | None = None) -> list[MoleculeInput]:
    rep = read_csv_report(path, smiles_col, id_col, smiles_preference)
    rep.emit(log)
    return rep.inputs


def _select_sheet(names: list[str], sheet: str | int | None, path: Path) -> str:
    if sheet is None or sheet == "":
        return names[0]
    if isinstance(sheet, str) and sheet.strip().isdigit() and sheet.strip().lower() not in {n.lower() for n in names}:
        sheet = int(sheet)           # "1" from the command line: a 0-based index unless a sheet is named "1"
    if isinstance(sheet, int):
        if 0 <= sheet < len(names):
            return names[sheet]
    else:
        lower = {n.lower(): n for n in names}
        if sheet.lower() in lower:
            return lower[sheet.lower()]
    raise ValueError(f"Sheet {sheet!r} not in {path.name} (sheets: {', '.join(names)})")


def _header_and_rows(values_iter: Iterable[Sequence]) -> tuple[list[str], list[tuple[int, int, dict]]]:
    """The header is the first non-empty row; returns (columns, [(row number, data index, row)])."""
    header: list[str] = []
    rows = []
    header_row = 0
    for rownum, values in enumerate(values_iter, 1):
        if not header:
            if any(_cell(v) for v in values):
                header, header_row = [_cell(v) for v in values], rownum
            continue
        rows.append((rownum, rownum - header_row, {h: v for h, v in zip(header, values) if h}))
    return [h for h in header if h], rows


def _excel_rows(path: Path, sheet: str | int | None) -> tuple[list[str], list[tuple[int, int, dict]]]:
    if path.suffix.lower() == ".xls":
        try:
            import xlrd
        except ImportError:
            raise ValueError(f"{path.name} is a legacy Excel 97-2003 .xls file. Open it in Excel and "
                             "'Save As' .xlsx or .csv, then load that file (or `pip install xlrd`).") from None
        book = xlrd.open_workbook(str(path))
        ws = book.sheet_by_name(_select_sheet(book.sheet_names(), sheet, path))
        return _header_and_rows(ws.row_values(r) for r in range(ws.nrows))
    try:
        import openpyxl
    except ImportError:
        raise ValueError("Reading .xlsx needs openpyxl (`pip install openpyxl`), or save the sheet as .csv") from None
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)  # data_only: cached formula values
    try:
        ws = wb[_select_sheet(wb.sheetnames, sheet, path)]
        return _header_and_rows(ws.iter_rows(values_only=True))
    finally:
        wb.close()  # read-only workbooks keep the file open (and locked on Windows) until closed


def read_excel_report(path: Path, smiles_col: str | None = None, id_col: str | None = None,
                      sheet: str | int | None = None,
                      smiles_preference: str | Sequence[str] | None = None) -> ReadReport:
    """Read .xlsx/.xlsm (or .xls when xlrd is installed) with the same column logic as CSV.

    The first sheet is used unless ``sheet`` (name or 0-based index) says otherwise; the header is the
    first non-empty row. CompTox Dashboard exports work as downloaded: DTXSID becomes the id,
    PREFERRED_NAME the name, and every column is kept in ``meta`` (plus ``smiles_column``, the column
    the structure came from, and ``row``).
    """
    columns, rows = _excel_rows(path, sheet)
    return _rows_to_inputs(columns, rows, path.name, path.stem, smiles_col, id_col, smiles_preference)


def read_excel(path: Path, smiles_col: str | None = None, id_col: str | None = None,
               sheet: str | int | None = None, smiles_preference: str | Sequence[str] | None = None,
               log: Callable[[str], None] | None = None) -> list[MoleculeInput]:
    rep = read_excel_report(path, smiles_col, id_col, sheet, smiles_preference)
    rep.emit(log)
    return rep.inputs


def read_table(path: str | Path, smiles_col: str | None = None, id_col: str | None = None,
               sheet: str | int | None = None,
               smiles_preference: str | Sequence[str] | None = None) -> ReadReport:
    """CSV or spreadsheet -> ``ReadReport`` (inputs, skipped rows, which SMILES columns were used)."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".csv":
        return read_csv_report(path, smiles_col, id_col, smiles_preference)
    if ext in TABLE_TYPES:
        return read_excel_report(path, smiles_col, id_col, sheet, smiles_preference)
    raise ValueError(f"Not a table file: {path.name} (expected {', '.join(sorted(TABLE_TYPES))})")


def read_file(path: str | Path, smiles_col: str | None = None, id_col: str | None = None,
              charge: int = 0, sheet: str | int | None = None,
              smiles_preference: str | Sequence[str] | None = None,
              log: Callable[[str], None] | None = None) -> list[MoleculeInput]:
    """Read any supported file. For tables, skipped rows and SMILES-column fallbacks are reported
    through ``log`` (or this module's logger when no callback is given)."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext in TABLE_TYPES:
        rep = read_table(path, smiles_col, id_col, sheet, smiles_preference)
        rep.emit(log)
        return rep.inputs
    if ext in (".smi", ".smiles", ".txt"):
        return parse_text(path.read_text(encoding="utf-8"), skip_errors=True, log=log, source=path.name)
    if ext == ".sdf":
        out = []
        for i, mol in enumerate(Chem.SDMolSupplier(str(path), removeHs=False), 1):
            if mol is None:
                continue
            props = mol.GetPropsAsDict()
            ident = (props.get("DTXSID") or props.get("ID") or (mol.GetProp("_Name") if mol.HasProp("_Name") else "")
                     or f"{path.stem}_{i}")
            out.append(MoleculeInput(id=safe_id(str(ident)), name=str(ident), source=f"{path.name}#{i}", mol=mol,
                                     smiles=Chem.MolToSmiles(Chem.RemoveHs(mol)),
                                     meta={k: str(v) for k, v in props.items()}))
        return out
    if ext == ".mol":
        mol = Chem.MolFromMolFile(str(path), removeHs=False)
    elif ext == ".pdb":
        mol = Chem.MolFromPDBFile(str(path), removeHs=False)
    elif ext == ".xyz":
        mol = Chem.MolFromXYZFile(str(path))
        if mol is not None:
            rdDetermineBonds.DetermineBonds(mol, charge=charge)
    else:
        raise ValueError(f"Unsupported file type {ext} (supported: {', '.join(sorted(FILE_TYPES))})")
    if mol is None:
        raise ValueError(f"Could not read a molecule from {path.name}")
    return [MoleculeInput(id=safe_id(path.stem), name=path.stem, source=path.name, mol=mol,
                          smiles=Chem.MolToSmiles(Chem.RemoveHs(mol)))]
